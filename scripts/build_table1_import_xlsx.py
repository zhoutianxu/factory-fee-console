from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT = Path("outputs/table1_import/表1_真实格式导入.xlsx")


headers = [
    "年月",
    "工厂代码",
    "一级部门",
    "工厂",
    "移动类型",
    "订单号",
    "责任人编码",
    "日期",
    "物料组",
    "物料编码",
    "物料描述",
    "品牌",
    "型号",
    "项目名",
    "工序机型",
    "工艺分类",
    "机型类别",
    "制式",
    "交货数量",
    "确认交货数量",
    "货币",
    "单价基本币",
    "汇率",
    "单价折合CNY",
    "凭证金额CNY",
]


orders = [
    "1000001723", "1000001405", "1000001406", "1000001407", "1000001413",
    "1000001414", "1000000840", "1000001509", "1000001807", "1000001510",
    "1000000843", "1000001410", "1000001411", "1000001411", "1000001412",
    "1000001443", "1000001488", "1000001488", "1000001488", "1000001488",
    "1000001443", "1000001443", "1000001443", "1000001443", "1000001445",
    "1000001445", "1000001451", "1000001207", "1000001210", "1000000893",
    "1000000893", "1000001451", "1000001210", "1000001451", "1000000893",
    "1000000893", "1000001207", "1000001207", "1000001207", "1000001451",
    "1000001211", "1000001408", "1000001408", "1000001409", "1000001212",
    "1000001212", "1000001487", "1000001449", "1000001487", "1000001449",
    "1000001449", "1000001220", "1000001220", "1000001220",
]

materials = [
    ("10056380", "Infinix_X6871.F1_机背膜_PH.256+12_美银", "Infinix", "X6871", 2, 3.00),
    ("10056381", "Infinix_X6871.F1_机中框_EC.256+12_美银", "Infinix", "X6871", 101, 3.00),
    ("10056382", "Infinix_X6871.F1_机中框_EC.256+12_美银", "Infinix", "X6871", 540, 3.00),
    ("10056383", "Infinix_X6871.F1_机背盖_EC.256+12_美银", "Infinix", "X6871", 120, 3.00),
    ("10057469", "Infinix_X6850.H1_摄像头_AF.256+8_灰悦", "Infinix", "X6850", 190, 3.00),
    ("10066385", "Infinix_X6860.A1_背盖盖_EC.256+8_美悦", "Infinix", "X6860", 4, 2.50),
    ("10066629", "Infinix_X6860.A1_支架盖_EC.256+8_美悦", "Infinix", "X6860", 159, 2.50),
    ("10060629", "Infinix_X6860.A1_支架盖_EC.256+8_美悦", "Infinix", "X6860", 100, 2.50),
    ("10060630", "Infinix_X6860.A1_保护盖_EC.256+8_美悦", "Infinix", "X6860", 1, 2.50),
    ("10060630", "Infinix_X6860.A1_保护盖_EC.256+8_美悦", "Infinix", "X6860", 182, 2.50),
    ("10062133", "Infinix_X6861.B1_壳体盖_SA.256+12_美悦", "Infinix", "X6861", 38, 2.50),
    ("10062134", "Infinix_X6861.B1_壳体盖_SA.256+12_美悦", "Infinix", "X6861", 255, 2.50),
    ("10062134", "Infinix_X6861.B1_壳体盖_SA.256+12_美悦", "Infinix", "X6861", 140, 2.50),
    ("10062135", "Infinix_X6861.B1_壳体盖_SA.256+12_美悦", "Infinix", "X6861", 760, 2.50),
    ("10066035", "TECNO_CM6_A1_星河幻梦_AFR.256+8_美悦", "Tecno", "CM6", 100, 2.30),
    ("10066906", "TECNO_CM6_A1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM6", 20, 2.30),
    ("10066906", "TECNO_CM6_A1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM6", 80, 2.30),
    ("10066906", "TECNO_CM6_A1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM6", 80, 2.30),
    ("10066906", "TECNO_CM6_A1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM6", 1060, 2.30),
    ("10066906", "TECNO_CM6_A1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM6", 1120, 2.30),
    ("10066906", "TECNO_CM6_A1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM6", 740, 2.30),
    ("10066906", "TECNO_CM6_A1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM6", 40, 2.30),
    ("10066906", "TECNO_CM6_A1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM6", 960, 2.30),
    ("10066909", "TECNO_CM6_A1_星河幻梦_AFR.256+8_美悦", "Tecno", "CM6", 1120, 2.30),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 1120, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 800, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 360, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 20, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 1120, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 30, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 28, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 145, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 15, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 960, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 12, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 510, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 60, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 20, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 10, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 1120, 2.60),
    ("10066932", "TECNO_CM5.E1_星河幻梦_AFR.256+8_映悦", "Tecno", "CM5", 780, 2.60),
    ("10066954", "Infinix_X6871.C1_机中框_MY.256+8_美悦", "Infinix", "X6871", 500, 3.00),
    ("10066954", "Infinix_X6871.C1_机中框_MY.256+8_美悦", "Infinix", "X6871", 960, 3.00),
    ("10066955", "Infinix_X6871.C1_机背盖_MY.256+8_美悦", "Infinix", "X6871", 500, 3.00),
    ("10068031", "TECNO_CM5.K1_星河幻梦_KH.256+8_映悦", "Tecno", "CM5", 500, 2.60),
    ("10068031", "TECNO_CM5.K1_星河幻梦_KH.256+8_映悦", "Tecno", "CM5", 1120, 2.60),
    ("10068031", "TECNO_CM5.K1_星河幻梦_KH.256+8_映悦", "Tecno", "CM5", 20, 2.60),
    ("10068116", "TECNO_CM6.A1_星色琳琅_MY.HOM.256+8_美悦", "Tecno", "CM6", 100, 2.30),
    ("10068116", "TECNO_CM6.A1_星色琳琅_MY.HOM.256+8_美悦", "Tecno", "CM6", 240, 2.30),
    ("10068116", "TECNO_CM6.A1_星色琳琅_MY.HOM.256+8_美悦", "Tecno", "CM6", 640, 2.30),
    ("10068116", "TECNO_CM6.A1_星色琳琅_MY.HOM.256+8_美悦", "Tecno", "CM6", 40, 2.30),
    ("10068124", "TECNO_CM6.D1_星河幻梦_KH.256+8_映悦", "Tecno", "CM6", 20, 2.30),
    ("10068124", "TECNO_CM6.D1_星河幻梦_KH.256+8_映悦", "Tecno", "CM6", 20, 2.30),
    ("10068124", "TECNO_CM6.D1_星河幻梦_KH.256+8_映悦", "Tecno", "CM6", 720, 2.30),
]


def build_rows():
    rows = []
    for idx, (material_code, description, brand, model, qty, unit_price) in enumerate(materials):
        rows.append([
            "202502",
            "1051",
            "生产中心",
            "零部件厂",
            "101",
            orders[idx],
            "1051",
            date(2025, 2, 28),
            "",
            material_code,
            description,
            brand,
            model,
            "",
            "整机",
            "组装",
            "智能机",
            "5G" if brand == "Infinix" or idx in {13, 14, 15, 41, 42, 43} else "4G",
            qty,
            qty,
            "CNY",
            unit_price,
            1.00,
            unit_price,
            round(qty * unit_price, 2),
        ])
    return rows


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append(headers)
    for row in build_rows():
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    text_cols = {1, 2, 5, 6, 7, 9, 10, 12, 13, 18, 21}
    numeric_cols = {19, 20, 22, 23, 24, 25}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if cell.column in text_cols:
                cell.number_format = "@"
            elif cell.column == 8:
                cell.number_format = "yyyy-mm-dd"
            elif cell.column in numeric_cols:
                cell.number_format = "0.00" if cell.column in {22, 23, 24, 25} else "0"

    widths = {
        "A": 10, "B": 10, "C": 12, "D": 12, "E": 10, "F": 14, "G": 12,
        "H": 12, "I": 10, "J": 12, "K": 46, "L": 12, "M": 10, "N": 10,
        "O": 12, "P": 12, "Q": 12, "R": 8, "S": 12, "T": 14, "U": 8,
        "V": 14, "W": 8, "X": 14, "Y": 14,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:Y{ws.max_row}"
    ws.sheet_view.showGridLines = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, data_only=True)
    sheet = check["Sheet1"]
    assert sheet.max_row == len(materials) + 1
    assert sheet.max_column == len(headers)
    assert sheet["A1"].value == "年月"
    assert sheet["Y2"].value == 6
    print(f"saved {OUTPUT} rows={sheet.max_row - 1} cols={sheet.max_column}")


if __name__ == "__main__":
    main()
